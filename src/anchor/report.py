"""Figures, README GIFs, results.xlsx and a metrics summary."""
from __future__ import annotations

import json
from pathlib import Path

import matplotlib
import matplotlib.ticker

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from matplotlib.animation import FuncAnimation, PillowWriter  # noqa: E402

from .data import MONOMERS, ROOT, SOURCES, weights  # noqa: E402
from .models import Anchor, cv_plus  # noqa: E402
from .physics import fox_tg  # noqa: E402

OUT = ROOT / "outputs"
FIG = OUT / "figures"
ASSETS = ROOT / "assets"

INK, MUTED, GRID = "#1E1B2E", "#6B6880", "#E6E4EC"
COLORS = {"Fox": "#8C8A96", "Pure ML": "#E08A1E", "ANCHOR": "#502B8B", "Random": "#B8B5C2"}
K0 = 273.15

plt.rcParams.update({
    "font.family": "DejaVu Sans", "font.size": 13, "axes.edgecolor": MUTED, "axes.labelcolor": INK,
    "xtick.color": MUTED, "ytick.color": MUTED, "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.8, "axes.titleweight": "bold",
    "axes.titlecolor": INK, "figure.dpi": 100, "savefig.dpi": 300, "savefig.bbox": "tight",
})


def _save(fig, name: str) -> Path:
    FIG.mkdir(parents=True, exist_ok=True)
    path = FIG / name
    fig.savefig(path, facecolor="white")
    plt.close(fig)
    return path


def parity(pred: pd.DataFrame, summary: dict) -> Path:
    y = pred["Tg_exp_K"] - K0
    lo, hi = y.min() - 15, y.max() + 15
    fig, axes = plt.subplots(1, 3, figsize=(15, 5.2), sharex=True, sharey=True)
    for ax, name in zip(axes, ("Fox", "Pure ML", "ANCHOR")):
        p = pred[f"{name}_K"] - K0
        ax.plot([lo, hi], [lo, hi], color=MUTED, lw=1, ls="--", zorder=1)
        if name != "Fox":
            err = np.vstack([p - (pred[f"{name}_lo_K"] - K0), (pred[f"{name}_hi_K"] - K0) - p])
            ax.errorbar(y, p, yerr=np.clip(err, 0, None), fmt="none", ecolor=COLORS[name], alpha=0.35, lw=1.2, zorder=2)
        ax.scatter(y, p, s=46, color=COLORS[name], edgecolor="white", linewidth=0.8, zorder=3)
        ax.set_title(f"{name}   MAE {summary[name]['MAE_K']:.1f} K")
        ax.set_xlim(lo, hi), ax.set_ylim(lo, hi)
        ax.set_xlabel("Measured Tg (°C)")
    axes[0].set_ylabel("Predicted Tg (°C)")
    fig.suptitle("Each monomer system predicted without seeing it (leave-one-system-out)", x=0.01, ha="left",
                 fontsize=15, fontweight="bold", color=INK)
    fig.subplots_adjust(top=0.84)
    return _save(fig, "parity.png")


def calibration(curves: dict) -> Path:
    fig, ax = plt.subplots(figsize=(6.2, 5.6))
    ax.plot([0.45, 1], [0.45, 1], color=MUTED, ls="--", lw=1, label="Ideal")
    for name in ("Pure ML", "ANCHOR"):
        lv = sorted(curves[name])
        ax.plot(lv, [curves[name][l] for l in lv], "o-", color=COLORS[name], lw=2.2, ms=7, label=name)
    ax.set_xlabel("Nominal interval confidence")
    ax.set_ylabel("Observed coverage (held-out systems)")
    ax.set_title("Are the uncertainty intervals honest?")
    ax.set_xlim(0.45, 1), ax.set_ylim(0, 1.02)
    ax.legend(frameon=False, loc="lower right")
    return _save(fig, "calibration.png")


def trials_figure(curves: pd.DataFrame, trials: pd.DataFrame, trial_summary: dict) -> Path:
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(14, 5.2), gridspec_kw={"width_ratios": [1.25, 1]})
    # Left: learning curves averaged over systems, only up to the step every system reaches.
    last_common = int(curves.groupby("system")["measured"].max().min())
    curves = curves[curves["measured"] <= last_common]
    for strategy, color in (("Random", COLORS["Random"]), ("ANCHOR (most uncertain)", COLORS["ANCHOR"])):
        g = curves[curves["strategy"] == strategy]
        per_seed = g.groupby(["measured", "seed"])["MAE_K"].mean().unstack()
        med = per_seed.median(axis=1)
        a1.plot(med.index, med.values, "o-", color=color, lw=2.4, label=strategy)
        if per_seed.shape[1] > 1:
            a1.fill_between(med.index, per_seed.quantile(0.25, axis=1), per_seed.quantile(0.75, axis=1),
                            color=color, alpha=0.25, lw=0)
    a1.set_xlabel("Rows of the new system measured")
    a1.xaxis.set_major_locator(matplotlib.ticker.MaxNLocator(integer=True))
    a1.set_ylabel("MAE on its unmeasured rows (K)")
    a1.set_title("Learning a new monomer system")
    a1.legend(frameon=False)
    # Right: trials to hit a target Tg within +/-5 K.
    order = ["Random", "Fox", "ANCHOR"]
    means = [trial_summary[s]["mean_trials"] for s in order]
    bars = a2.bar(order, means, color=[COLORS[s] for s in order], width=0.6)
    for b, s in zip(bars, order):
        a2.text(b.get_x() + b.get_width() / 2, b.get_height() + 0.05, f"{trial_summary[s]['mean_trials']:.2f}",
                ha="center", va="bottom", fontweight="bold", color=INK)
    a2.set_ylabel("Mean trials to hit target Tg ± 5 K")
    a2.set_title("Trials to reach a target")
    a2.grid(axis="x", visible=False)
    return _save(fig, "trials.png")


def residuals_figure(sys_table: pd.DataFrame, gt: pd.DataFrame) -> Path:
    t = sys_table.merge(gt, on="system").sort_values("Fox_bias_K")
    fig, ax = plt.subplots(figsize=(10, 5))
    colors = [COLORS["ANCHOR"] if v > 0 else COLORS["Pure ML"] for v in t["Fox_bias_K"]]
    ax.barh(t["system"], t["Fox_bias_K"], color=colors)
    ax.axvline(0, color=MUTED, lw=1)
    right = t["Fox_bias_K"].max() + 2
    for i, (_, r) in enumerate(t.iterrows()):
        a, b = r["system"].split("/")
        label = ("k ill-defined (homopolymer Tgs nearly equal)" if abs(r["k_fox_equivalent"] - 1) < 0.02
                 else f"Gordon–Taylor k = {r['k_fitted']:.2f}  (Fox ≡ {r['k_fox_equivalent']:.2f})")
        ax.text(right, i, label, va="center", ha="left", fontsize=11, color=MUTED)
    ax.set_xlim(t["Fox_bias_K"].min() - 1.5, right + 13)
    ax.set_xlabel("Mean measured − Fox Tg (K)")
    ax.set_title("Where physics alone misses: the gap ANCHOR learns")
    ax.grid(axis="y", visible=False)
    return _save(fig, "residuals_by_system.png")


def deck_figures(pred: pd.DataFrame, summary: dict, curves: dict, trial_summary: dict) -> list[Path]:
    """Three figures sized for the 3.69 x 1.79 in boxes on slide 5 (drawn at 2x)."""
    big = {"font.size": 19, "axes.titlesize": 20, "axes.labelsize": 18, "xtick.labelsize": 16, "ytick.labelsize": 16,
           "legend.fontsize": 16}
    size = (7.38, 3.58)
    paths = []
    with plt.rc_context(big):
        fig, ax = plt.subplots(figsize=size)
        y = pred["Tg_exp_K"] - K0
        lo, hi = y.min() - 15, y.max() + 15
        ax.plot([lo, hi], [lo, hi], color=MUTED, ls="--", lw=1.2)
        for name, mk in (("Fox", "s"), ("ANCHOR", "o")):
            ax.scatter(y, pred[f"{name}_K"] - K0, s=55, marker=mk, color=COLORS[name], edgecolor="white", lw=0.7,
                       label=f"{name} · MAE {summary[name]['MAE_K']:.1f} K", zorder=3 if name == "ANCHOR" else 2)
        ax.set_xlim(lo, hi), ax.set_ylim(lo, hi)
        ax.set_xlabel("Measured Tg (°C)"), ax.set_ylabel("Predicted (°C)")
        ax.legend(frameon=False, loc="upper left", handletextpad=0.2)
        paths.append(_save(fig, "deck_parity.png"))

        fig, ax = plt.subplots(figsize=size)
        ax.plot([0.45, 1], [0.45, 1], color=MUTED, ls="--", lw=1.2, label="Ideal")
        for name in ("Pure ML", "ANCHOR"):
            lv = sorted(curves[name])
            ax.plot(lv, [curves[name][l] for l in lv], "o-", color=COLORS[name], lw=2.6, ms=8, label=name)
        ax.set_xlim(0.45, 1), ax.set_ylim(0, 1.02)
        ax.set_xlabel("Nominal confidence"), ax.set_ylabel("Observed")
        ax.legend(frameon=False, loc="lower right")
        paths.append(_save(fig, "deck_calibration.png"))

        fig, ax = plt.subplots(figsize=size)
        order = ["Random", "Fox", "ANCHOR"]
        vals = [trial_summary[s]["mean_trials"] for s in order]
        bars = ax.barh(order, vals, color=[COLORS[s] for s in order], height=0.62)
        for b, v in zip(bars, vals):
            ax.text(v + 0.04, b.get_y() + b.get_height() / 2, f"{v:.2f}", va="center", fontweight="bold", color=INK)
        ax.set_xlabel(f"Mean trials to hit target Tg ± {5:.0f} K")
        ax.set_xlim(0, max(vals) * 1.25)
        ax.grid(axis="y", visible=False)
        paths.append(_save(fig, "deck_trials.png"))
    return paths


# ---------------------------------------------------------------- GIFs

def gif_hold_out(pred: pd.DataFrame) -> Path:
    """Systems are held out one by one; Fox vs ANCHOR predictions appear."""
    systems = list(pred.groupby("system")["Tg_exp_K"].mean().sort_values().index)
    y_all = pred["Tg_exp_K"] - K0
    lo, hi = y_all.min() - 15, y_all.max() + 15
    fig, axes = plt.subplots(1, 2, figsize=(10, 4.9), sharex=True, sharey=True)

    def draw(i):
        shown = pred[pred["system"].isin(systems[: i + 1])]
        current = pred[pred["system"] == systems[min(i, len(systems) - 1)]]
        for ax, name in zip(axes, ("Fox", "ANCHOR")):
            ax.clear()
            ax.plot([lo, hi], [lo, hi], color=MUTED, ls="--", lw=1)
            ax.scatter(shown["Tg_exp_K"] - K0, shown[f"{name}_K"] - K0, s=42, color=COLORS[name], alpha=0.35,
                       edgecolor="white", lw=0.6)
            if i < len(systems):
                ax.scatter(current["Tg_exp_K"] - K0, current[f"{name}_K"] - K0, s=90, color=COLORS[name],
                           edgecolor=INK, lw=1.2, zorder=5)
            mae = np.mean(np.abs(shown[f"{name}_K"] - shown["Tg_exp_K"]))
            ax.set_title(f"{name}   MAE so far {mae:.1f} K", fontsize=13)
            ax.set_xlim(lo, hi), ax.set_ylim(lo, hi)
            ax.set_xlabel("Measured Tg (°C)")
        axes[0].set_ylabel("Predicted Tg (°C)")
        label = f"Held out: {systems[i]}" if i < len(systems) else "All systems held out"
        fig.suptitle(f"Leave-one-system-out   ·   {label}", x=0.02, ha="left", fontweight="bold", color=INK)

    frames = list(range(len(systems))) + [len(systems)] * 3
    anim = FuncAnimation(fig, lambda f: draw(min(f, len(systems))), frames=frames)
    ASSETS.mkdir(exist_ok=True)
    path = ASSETS / "anchor_hold_out.gif"
    anim.save(path, writer=PillowWriter(fps=1), dpi=90)
    plt.close(fig)
    return path


def gif_learning(df: pd.DataFrame, tg_homo: np.ndarray, system: str, confidence: float = 0.9) -> Path:
    """One held-out binary system: ANCHOR asks for the most uncertain sample next."""
    a, b = system.split("/")
    base, pool = df[df["system"] != system], df[df["system"] == system].reset_index(drop=True)
    w_b, y_b = weights(base), base["Tg_exp_K"].to_numpy()
    w_p, y_p = weights(pool), pool["Tg_exp_K"].to_numpy()
    xa = pool[f"{a}_wt_frac"].to_numpy()

    grid = np.linspace(0, 1, 101)
    w_grid = np.zeros((len(grid), len(MONOMERS)))
    w_grid[:, MONOMERS.index(a)] = grid
    w_grid[:, MONOMERS.index(b)] = 1 - grid
    fox_grid = fox_tg(w_grid, tg_homo) - K0

    states, taken = [], []
    for _ in range(len(pool)):
        w_tr = np.vstack([w_b, w_p[taken]]) if taken else w_b
        y_tr = np.concatenate([y_b, y_p[taken]]) if taken else y_b
        pred = cv_plus(Anchor, tg_homo, w_tr, y_tr, np.vstack([w_grid, w_p]))
        lo, hi = pred.interval(confidence)
        n = len(grid)
        sd_pool = pred.std[n:]
        rest = [i for i in range(len(pool)) if i not in taken]
        nxt = rest[int(np.argmax(sd_pool[rest]))] if rest else None
        mae = float(np.mean(np.abs(pred.mean[n:][rest] - y_p[rest]))) if rest else float("nan")
        states.append((list(taken), pred.mean[:n] - K0, lo[:n] - K0, hi[:n] - K0, nxt, mae))
        if nxt is None:
            break
        taken.append(nxt)

    fig, ax = plt.subplots(figsize=(9, 5.2))
    ylo, yhi = min(y_p.min() - K0, fox_grid.min()) - 25, max(y_p.max() - K0, fox_grid.max()) + 25

    def draw(i):
        tk, mu, lo, hi, nxt, mae = states[i]
        ax.clear()
        ax.fill_between(grid, lo, hi, color=COLORS["ANCHOR"], alpha=0.16, lw=0, label="ANCHOR 90% interval")
        ax.plot(grid, fox_grid, color=COLORS["Fox"], lw=2, ls="--", label="Fox (physics only)")
        ax.plot(grid, mu, color=COLORS["ANCHOR"], lw=2.6, label="ANCHOR")
        unseen = [j for j in range(len(pool)) if j not in tk]
        ax.scatter(xa[unseen], y_p[unseen] - K0, s=70, facecolor="white", edgecolor=MUTED, lw=1.4,
                   label="Not yet measured", zorder=4)
        if tk:
            ax.scatter(xa[tk], y_p[tk] - K0, s=80, color=INK, zorder=5, label="Measured")
        if nxt is not None:
            ax.scatter([xa[nxt]], [y_p[nxt] - K0], s=320, facecolor="none", edgecolor=COLORS["Pure ML"], lw=2.6,
                       zorder=6, label="Next: most uncertain")
        ax.set_xlim(0, 1), ax.set_ylim(ylo, yhi)
        ax.set_xlabel(f"{a} weight fraction in {a}/{b} copolymer")
        ax.set_ylabel("Tg (°C)")
        status = f"MAE on unmeasured rows {mae:.1f} K" if not np.isnan(mae) else "all rows measured"
        ax.set_title(f"Experiments run: {len(tk)}   ·   {status}", loc="left", fontsize=13)
        ax.legend(frameon=True, framealpha=0.92, edgecolor=GRID, fontsize=10, loc="upper right")

    frames = list(range(len(states))) + [len(states) - 1] * 2
    anim = FuncAnimation(fig, draw, frames=frames)
    path = ASSETS / "anchor_active_learning.gif"
    anim.save(path, writer=PillowWriter(fps=0.9), dpi=90)
    plt.close(fig)
    return path


# ---------------------------------------------------------------- tables

def write_excel(path: Path, sheets: dict[str, pd.DataFrame], readme_rows: list[tuple[str, str]]) -> Path:
    from openpyxl.styles import Alignment, Font, PatternFill

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame(readme_rows, columns=["Item", "Description"]).to_excel(xw, sheet_name="README", index=False)
        for name, frame in sheets.items():
            frame.to_excel(xw, sheet_name=name[:31], index=False)
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="502B8B")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 70)
            ws.freeze_panes = "A2"
    return path


def write_metrics_md(path: Path, metrics: dict) -> Path:
    s = metrics["held_out"]
    t = metrics["trials_to_target"]
    lines = [
        "# ANCHOR headline numbers",
        "",
        f"Data: {metrics['data']['n_rows']} measured copolymers, {metrics['data']['n_systems']} monomer systems, "
        f"{metrics['data']['n_sources']} sources. Validation: leave-one-monomer-system-out.",
        "",
        "| Model | MAE (K) | RMSE (K) | R² | 90% coverage | Mean width (K) |",
        "| --- | --- | --- | --- | --- | --- |",
    ]
    for name in ("Fox", "Pure ML", "ANCHOR"):
        m = s[name]
        cov = f"{m['coverage_90']:.0%}" if "coverage_90" in m else "—"
        wid = f"{m['mean_width_K']:.1f}" if "mean_width_K" in m else "—"
        lines.append(f"| {name} | {m['MAE_K']:.1f} | {m['RMSE_K']:.1f} | {m['R2']:.2f} | {cov} | {wid} |")
    lines += [
        "",
        f"Trials to hit a target Tg within ±{metrics['tolerance_K']:.0f} K (mean over every row as target): "
        f"Random {t['Random']['mean_trials']:.2f}, Fox {t['Fox']['mean_trials']:.2f}, ANCHOR {t['ANCHOR']['mean_trials']:.2f}.",
        f"First-try hit rate: Random {t['Random']['first_try_hit_rate']:.0%}, Fox {t['Fox']['first_try_hit_rate']:.0%}, "
        f"ANCHOR {t['ANCHOR']['first_try_hit_rate']:.0%}.",
        "",
        f"External stress test (12 excluded Taheri resins, AA omitted, low Mn): Fox MAE "
        f"{metrics['external']['Fox_MAE_K']:.1f} K, ANCHOR MAE {metrics['external']['ANCHOR_MAE_K']:.1f} K, "
        f"ANCHOR 90% coverage {metrics['external']['ANCHOR_coverage_90']:.0%}.",
    ]
    path.write_text("\n".join(lines) + "\n")
    return path


def write_json(path: Path, obj: dict) -> Path:
    path.write_text(json.dumps(obj, indent=2) + "\n")
    return path


def sources_table() -> pd.DataFrame:
    return pd.DataFrame(sorted(SOURCES.items()), columns=["source_id", "citation"])
